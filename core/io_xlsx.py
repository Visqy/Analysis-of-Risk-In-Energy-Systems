"""Ekspor laporan dan impor parameter dalam format xlsx.

Ekspor (`build_export_workbook`) membuat 6 sheet — Ringkasan, Parameter, Hasil
Simulasi, Perbandingan Skenario, Monte Carlo, Info Versi — beserta grafik native
Excel. Impor (`parse_import_workbook`) membaca sheet "Parameter" dengan
validasi semua-atau-tidak-sama-sekali.
"""

import io

import numpy as np
import pandas as pd
from openpyxl.chart import BarChart, Reference, ScatterChart, Series

from core.inputs import (
    ALL_KEYS, DIST_OPTIONS, FIELD_SPECS, SOURCE_OPTIONS, fmt_input, parse_field, scenario_names,
)
from core.scenarios import PARAM_DISCLAIMER, VARIABLE_UNITS

APP_VERSION = "ARIES v0.1 – Prototipe Penelitian"


def _config_values(cfg):
    """Nilai input (sesuai satuan form) dari config run."""
    v = {"E0": cfg["y0"][0], "D0": cfg["y0"][1], "P0": cfg["y0"][2], "T_end": cfg["T_end"],
         "n_sim": cfg["n_sim"], "conf": int(round(cfg["conf"] * 100)), "seed": cfg["seed"]}
    v.update(cfg["params"])
    v.update({f"sd_{k}": s for k, s in cfg["sd"].items()})
    return v


def _scatter(ws, title, x_col, y_cols, n_rows, anchor, y_title):
    ch = ScatterChart()
    ch.title = title
    ch.style = 13
    ch.x_axis.title = "t (tahun)"
    ch.y_axis.title = y_title
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.height, ch.width = 7.5, 15
    xref = Reference(ws, min_col=x_col, min_row=2, max_row=n_rows + 1)
    for c in y_cols:
        s = Series(Reference(ws, min_col=c, min_row=1, max_row=n_rows + 1), xref,
                   title_from_data=True)
        s.marker.symbol = "none"
        ch.series.append(s)
    ws.add_chart(ch, anchor)


def _bar(ws, title, cats_col, val_col, n_rows, anchor, y_title, horizontal=False, extra_cols=()):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = title
    ch.y_axis.title = y_title
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.height, ch.width = 7.5, 15
    for c in (val_col, *extra_cols):
        ch.add_data(Reference(ws, min_col=c, min_row=1, max_row=n_rows + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cats_col, min_row=2, max_row=n_rows + 1))
    ws.add_chart(ch, anchor)


def build_export_workbook(run):
    """Bangun laporan xlsx dari snapshot `run` dan kembalikan isinya (bytes).

    Memuat input, asumsi, metadata (seed, tanggal, versi), hasil simulasi,
    perbandingan skenario (peringkat berdasarkan total kerugian dan berdasarkan risiko/CVaR), ringkasan Monte Carlo (interval kuantil 95%), dan
    grafik: dinamika E/D/P/L, total kerugian per skenario, histogram, tornado.
    """
    cfg = run["config"]
    src = cfg["source"]
    meta = run["meta"]
    losses = run["losses"]
    q_low, q_high = np.percentile(losses, [2.5, 97.5])
    conf_pct = int(round(cfg["conf"] * 100))
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary = [
            ("Sistem", APP_VERSION),
            ("Catatan", PARAM_DISCLAIMER),
            ("Versi model", meta["model_version"]),
            ("Tanggal analisis", meta["date"]),
            ("Skenario aktif", cfg["scenario"]),
            ("Periode analisis (tahun)", cfg["T_end"]),
            ("Jumlah iterasi", cfg["n_sim"]),
            ("Seed acak", cfg["seed"]),
            ("Distribusi ketidakpastian", DIST_OPTIONS[cfg["dist"]]),
            ("Tingkat keyakinan VaR/CVaR (%)", conf_pct),
            ("Status sumber data", src["status"]),
            ("Sumber data", src["name"]),
            ("Tahun data", src["year"]),
            (f"VaR {conf_pct}% ({VARIABLE_UNITS['L']})", run["var"]),
            (f"CVaR {conf_pct}% ({VARIABLE_UNITS['L']})", run["cvar"]),
            (f"Interval kuantil 95% bawah (persentil 2,5)", float(q_low)),
            (f"Interval kuantil 95% atas (persentil 97,5)", float(q_high)),
            (f"E(T) ({VARIABLE_UNITS['E']})", float(run["E"][-1])),
            (f"D(T) ({VARIABLE_UNITS['D']})", float(run["D"][-1])),
            (f"P(T) ({VARIABLE_UNITS['P']})", float(run["P"][-1])),
            (f"Total kerugian kumulatif ({VARIABLE_UNITS['L']})", run["total_loss"]),
        ]
        pd.DataFrame(summary, columns=["Item", "Nilai"]).to_excel(
            writer, sheet_name="Ringkasan", index=False)

        values = _config_values(cfg)
        rows = [{"Key": k, "Parameter": FIELD_SPECS[k]["label"], "Nilai": values[k],
                 "Satuan": FIELD_SPECS[k]["unit"], "Min": FIELD_SPECS[k]["min"],
                 "Max": FIELD_SPECS[k]["max"]} for k in ALL_KEYS]
        extras = {"scenario_name": cfg["scenario"],
                  "source_status": src["status"], "source_name": src["name"],
                  "source_year": src["year"]}
        rows += [{"Key": k, "Parameter": k, "Nilai": v, "Satuan": "", "Min": "", "Max": ""}
                 for k, v in extras.items()]
        pd.DataFrame(rows).to_excel(writer, sheet_name="Parameter", index=False)

        n = len(run["t"])
        pd.DataFrame({
            "t (tahun)": run["t"],
            f"E(t) ({VARIABLE_UNITS['E']})": run["E"],
            f"D(t) ({VARIABLE_UNITS['D']})": run["D"],
            f"P(t) ({VARIABLE_UNITS['P']})": run["P"],
            f"L(t) ({VARIABLE_UNITS['L']})": run["L"],
        }).to_excel(writer, sheet_name="Hasil Simulasi", index=False)
        ws = writer.sheets["Hasil Simulasi"]
        _scatter(ws, "Kapasitas E(t) dan Permintaan D(t)", 1, [2, 3], n, "H2", VARIABLE_UNITS["E"])
        _scatter(ws, "Harga P(t)", 1, [4], n, "H18", VARIABLE_UNITS["P"])
        _scatter(ws, "Kerugian ekonomi L(t)", 1, [5], n, "H34", VARIABLE_UNITS["L"])

        rank_rows = run["ranking"]
        risk = {r["Skenario"]: r for r in run["ranking_risk"]}
        pd.DataFrame([
            {"Peringkat": r["rank"], "Skenario": r["Skenario"],
             f"Total Kerugian ({VARIABLE_UNITS['L']})": float(r["loss"]),
             "Penurunan (absolut)": float(r["delta"]), "Penurunan (%)": float(r["pct"]),
             f"Peringkat Risiko (CVaR {conf_pct}%)": risk[r["Skenario"]]["rank"],
             f"CVaR {conf_pct}%": float(risk[r["Skenario"]]["cvar"]),
             f"VaR {conf_pct}%": float(risk[r["Skenario"]]["var"])}
            for r in rank_rows
        ]).to_excel(writer, sheet_name="Perbandingan Skenario", index=False)
        ws = writer.sheets["Perbandingan Skenario"]
        _bar(ws, "Total kerugian per skenario", 2, 3, len(rank_rows), "J2", VARIABLE_UNITS["L"])

        stats = pd.DataFrame([
            ("Jumlah iterasi", cfg["n_sim"]), ("Seed acak", cfg["seed"]),
            ("Rata-rata", float(losses.mean())), ("Median", float(np.median(losses))),
            ("Simpangan baku", float(losses.std())),
            (f"VaR {conf_pct}%", run["var"]), (f"CVaR {conf_pct}%", run["cvar"]),
            ("Interval kuantil 95% (persentil 2,5)", float(q_low)),
            ("Interval kuantil 95% (persentil 97,5)", float(q_high)),
        ], columns=["Statistik", f"Nilai ({VARIABLE_UNITS['L']})"])
        stats.to_excel(writer, sheet_name="Monte Carlo", index=False)
        ws = writer.sheets["Monte Carlo"]

        unc_rows = []
        for name, spec in run["uncertain"].items():
            unc_rows.append({"Parameter": name, "Distribusi": spec[0], "Mean": spec[1],
                             "Std": spec[2],
                             "Batas bawah": spec[3] if spec[0] == "truncnormal" else 1e-6,
                             "Batas atas": spec[4] if spec[0] == "truncnormal" else ""})
        pd.DataFrame(unc_rows).to_excel(writer, sheet_name="Monte Carlo", index=False,
                                        startrow=len(stats) + 2)

        counts, edges = np.histogram(losses, bins=20)
        hist = pd.DataFrame({"Kelas (tengah)": np.round((edges[:-1] + edges[1:]) / 2, 2),
                             "Frekuensi": counts})
        hist_row = len(stats) + 2 + len(unc_rows) + 3
        hist.to_excel(writer, sheet_name="Monte Carlo", index=False, startrow=hist_row)

        sens = pd.DataFrame([
            {"Parameter": s[0], "-20%": s[1] - s[2], "+20%": s[3] - s[2]} for s in run["sens"]
        ])
        sens_row = hist_row + len(hist) + 3
        sens.to_excel(writer, sheet_name="Monte Carlo", index=False, startrow=sens_row)

        h = Reference(ws, min_col=2, min_row=hist_row + 1, max_row=hist_row + len(hist) + 1)
        hc = BarChart()
        hc.type = "col"
        hc.title = "Distribusi total kerugian (histogram)"
        hc.y_axis.title = "Frekuensi"
        hc.x_axis.title = f"Total kerugian ({VARIABLE_UNITS['L']})"
        hc.x_axis.delete = False
        hc.y_axis.delete = False
        hc.add_data(h, titles_from_data=True)
        hc.set_categories(Reference(ws, min_col=1, min_row=hist_row + 2,
                                    max_row=hist_row + len(hist) + 1))
        hc.height, hc.width = 7.5, 15
        ws.add_chart(hc, "H2")

        tc = BarChart()
        tc.type = "bar"
        tc.title = "Sensitivitas parameter (±20%)"
        tc.x_axis.delete = False
        tc.y_axis.delete = False
        tc.add_data(Reference(ws, min_col=2, max_col=3, min_row=sens_row + 1,
                              max_row=sens_row + len(sens) + 1), titles_from_data=True)
        tc.set_categories(Reference(ws, min_col=1, min_row=sens_row + 2,
                                    max_row=sens_row + len(sens) + 1))
        tc.height, tc.width = 7.5, 15
        ws.add_chart(tc, "H18")

        pd.DataFrame({"Info": [APP_VERSION, PARAM_DISCLAIMER]}).to_excel(
            writer, sheet_name="Info Versi", index=False)

    return buffer.getvalue()


def parse_import_workbook(file):
    """Baca sheet "Parameter" dari file xlsx hasil Unduh Laporan.

    Semua field input wajib ada dan valid (rentang, bilangan bulat, tidak
    kosong); bila ada yang salah tidak ada yang diterapkan.

    Returns:
        (updates, errors): `updates` = dict kunci-state -> nilai siap untuk
        `state.request_set` (kosong bila ada error); `errors` = daftar pesan.
    """
    try:
        df = pd.read_excel(file, sheet_name="Parameter")
    except Exception as exc:
        return {}, [f"Gagal membaca file: {exc}"]
    if "Key" not in df.columns or "Nilai" not in df.columns:
        return {}, ["Sheet 'Parameter' harus punya kolom 'Key' dan 'Nilai' "
                    "(format hasil Unduh Laporan)."]

    cells = {str(r["Key"]): r["Nilai"] for _, r in df.iterrows() if not pd.isna(r["Key"])}
    updates, errors = {}, []
    for k in ALL_KEYS:
        label = FIELD_SPECS[k]["label"]
        if k not in cells:
            errors.append(f"{label}: tidak ditemukan di file (file mungkin dari versi lama).")
            continue
        raw = cells[k]
        if pd.isna(raw):
            errors.append(f"{label}: nilai kosong.")
            continue
        value, err = parse_field(k, str(raw))
        if err:
            errors.append(f"{label}: {err}")
        else:
            updates[f"in_{k}"] = fmt_input(value, FIELD_SPECS[k]["dec"])

    def txt(k):
        """Nilai sel `k` sebagai teks ("" bila kosong)."""
        v = cells.get(k)
        return "" if v is None or pd.isna(v) else str(v)

    if errors:
        return {}, errors
    sc = txt("scenario_name")
    if sc in scenario_names():
        updates["scenario_name"] = sc
    if txt("source_status") in SOURCE_OPTIONS:
        updates["source_status"] = txt("source_status")
        updates["source_name"] = txt("source_name")
        yr = txt("source_year")
        updates["source_year"] = yr[:-2] if yr.endswith(".0") else yr
    return updates, []
