import io

import openpyxl
import pytest

from core.inputs import ALL_KEYS, default_text
from core.io_xlsx import build_export_workbook, parse_import_workbook

SHEETS = ["Ringkasan", "Parameter", "Hasil Simulasi", "Perbandingan Skenario",
          "Monte Carlo", "Info Versi"]


@pytest.fixture(scope="module")
def xlsx_bytes(small_run):
    return build_export_workbook(small_run)


def _edit(xlsx_bytes, fn):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _set_cell(key, value):
    def fn(wb):
        for row in wb["Parameter"].iter_rows(min_row=2):
            if row[0].value == key:
                row[2].value = value
    return fn


def test_sheet_dan_chart(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == SHEETS
    assert sum(len(ws._charts) for ws in wb) == 6


def test_ringkasan_memuat_metadata(xlsx_bytes):
    ws = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))["Ringkasan"]
    items = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2)}
    assert items["Versi model"] == "ARIES v0.1"
    assert items["Seed acak"] == 2026
    assert items["Jumlah iterasi"] == 100
    assert "Indeks risiko" not in items
    assert any("Interval kuantil" in k for k in items)
    assert not any("kepercayaan" in k.lower() for k in items)


def test_round_trip(xlsx_bytes):
    updates, errors = parse_import_workbook(io.BytesIO(xlsx_bytes))
    assert errors == []
    for k in ALL_KEYS:
        assert updates[f"in_{k}"] == default_text(k) or k == "n_sim"
    assert updates["in_n_sim"] == "100"
    assert updates["scenario_name"] == "Baseline (Business as Usual)"
    assert "dist_kind" not in updates


def test_tolak_di_luar_rentang(xlsx_bytes):
    updates, errors = parse_import_workbook(_edit(xlsx_bytes, _set_cell("E0", 9999)))
    assert updates == {} and any("Di luar rentang" in e for e in errors)


def test_tolak_nilai_kosong(xlsx_bytes):
    updates, errors = parse_import_workbook(_edit(xlsx_bytes, _set_cell("kappa", None)))
    assert updates == {} and any("kosong" in e for e in errors)


def test_all_or_nothing_banyak_error(xlsx_bytes):
    def fn(wb):
        _set_cell("E0", 9999)(wb)
        _set_cell("seed", "abc")(wb)
    updates, errors = parse_import_workbook(_edit(xlsx_bytes, fn))
    assert updates == {} and len(errors) == 2


def test_tolak_tanpa_kolom_key(xlsx_bytes):
    def fn(wb):
        wb["Parameter"]["A1"].value = "Nama"
    updates, errors = parse_import_workbook(_edit(xlsx_bytes, fn))
    assert updates == {} and "Key" in errors[0]


def test_tolak_file_lama_tanpa_kunci_baru(xlsx_bytes):
    def fn(wb):
        ws = wb["Parameter"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value in ("seed", "n_sim", "E0"):
                row[0].value = None
    updates, errors = parse_import_workbook(_edit(xlsx_bytes, fn))
    assert updates == {} and any("versi lama" in e for e in errors)


def test_bukan_xlsx():
    updates, errors = parse_import_workbook(io.BytesIO(b"bukan excel"))
    assert updates == {} and errors


def test_sumber_data_di_ringkasan_dan_round_trip():
    from core.run import run_simulation
    from tests.conftest import make_config
    cfg = make_config(n_sim=100)
    cfg["source"] = {"status": "Data Aktual", "name": "Statistik PLN", "year": "2024"}
    data = build_export_workbook(run_simulation(cfg))
    ws = openpyxl.load_workbook(io.BytesIO(data))["Ringkasan"]
    items = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2)}
    assert (items["Status sumber data"], items["Sumber data"], items["Tahun data"]) ==         ("Data Aktual", "Statistik PLN", "2024")
    updates, errors = parse_import_workbook(io.BytesIO(data))
    assert errors == []
    assert updates["source_status"] == "Data Aktual"
    assert updates["source_name"] == "Statistik PLN" and updates["source_year"] == "2024"


def test_perbandingan_memuat_peringkat_risiko(xlsx_bytes):
    ws = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))["Perbandingan Skenario"]
    header = [c.value for c in ws[1]]
    assert "Peringkat Risiko (CVaR 95%)" in header and "CVaR 95%" in header and "VaR 95%" in header
    rank_idx = header.index("Peringkat Risiko (CVaR 95%)")
    ranks = sorted(row[rank_idx].value for row in ws.iter_rows(min_row=2, max_row=6))
    assert ranks == [1, 2, 3, 4, 5]
